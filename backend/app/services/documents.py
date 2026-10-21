import io
import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader

ALLOWED_CV_EXTENSIONS = {".pdf", ".docx", ".txt"}
ALLOWED_TRANSCRIPT_EXTENSIONS = {".pdf", ".docx", ".txt", ".vtt", ".srt", ".md"}


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif ext == ".docx":
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    else:
        text = data.decode("utf-8", errors="ignore")
        if ext in {".vtt", ".srt"}:
            text = _strip_captions(text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _strip_captions(text: str) -> str:
    lines = []
    for line in text.splitlines():
        line = line.strip()
        if not line or line == "WEBVTT" or line.isdigit() or "-->" in line:
            continue
        lines.append(line)
    return "\n".join(lines)
